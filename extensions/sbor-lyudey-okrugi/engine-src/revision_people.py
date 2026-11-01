"""
Ревизские сказки: от строк таблицы к людям, родству и Excel.

Вход : то же, что у extract_rows.py (данные страниц __NEXT_DATA__: .json / .html / папка),
       либо уже готовый pages.json от extract_rows.py.
Выход: Excel (.xlsx) с листами «Люди», «Семьи», «Блоки текста», «Страницы», «Сводка», «Об источнике»;
       рядом JSON с теми же людьми (для своих скриптов и родословного древа).

Распознавание сканов НЕ выполняется, сеть не используется: только готовые тексты и координаты.

Что добавляется к строкам extract_rows.py (опыт обработки ф. 51, оп. 8, д. 588, сказка 1850 г. по имению
Н. А. Бутурлина, сканы 130–225):
  * текст Основного блока делится на людей (split_names.py) и сопоставляется со строками по порядку;
  * у каждого человека: имя в современной орфографии, отчество (если в тексте нет — от имени отца),
    отношение («глава», «сын Никиты Григорьева», «жена Степана Никитина», «второй сын …»)
    и ID отца/мужа — по нему строится древо;
  * возраст по прошлой ревизии и нынешний: «новорожд.», «½», «мес.», год рождения ≈;
  * выбытие: «умер 1847», «отдан в рекруты 1845», «переведён …», статус налицо / выбыл / в отлучке;
  * селение и сословие (крестьяне / дворовые) по шапке страницы с переносом на следующие сканы,
    номер семьи с переносом с предыдущего скана;
  * «?» у сомнительных значений (ячейка у границы, несколько значений в одном месте).
"""
import argparse
import json
import re
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract_rows as ER  # noqa: E402
from split_names import split_block, nk, names, RELATION, PRON  # noqa: E402

REVISIONS = {1: 1719, 2: 1744, 3: 1762, 4: 1782, 5: 1795, 6: 1811, 7: 1815, 8: 1834, 9: 1850, 10: 1858}
ROMAN = {8: "8-й", 9: "9-й", 10: "10-й"}
DOUBT_FLAGS = ("y_near", "weak_col", "multi_cell", "span_conflict")

# ----------------------------------------------------------------------------- орфография
_OLD = str.maketrans({"ѣ": "е", "Ѣ": "Е", "і": "и", "І": "И", "ѳ": "ф", "Ѳ": "Ф", "ѵ": "и", "Ѵ": "И",
                      "ѡ": "о", "ꙋ": "у", "ї": "и"})
_LAT = str.maketrans({"i": "и", "I": "И"})


def modern(text):
    """современная орфография: без ѣ, і, ѳ, ѵ и конечного ъ; латинская i в русском слове -> и"""
    out = []
    for w in (text or "").split():
        if re.search(r"[а-яА-ЯѣѢіІѳѲ]", w):
            w = w.translate(_LAT)
        w = w.replace("ѯ", "кс").translate(_OLD)
        w = re.sub(r"[ъЪ](?=[^а-яА-Яa-zA-Z]*$)", "", w)
        out.append(w)
    return " ".join(out)


def word(w):
    """слово имени: современная орфография, без знаков, с заглавной"""
    w = re.sub(r"[^\w-]", "", modern(w))
    return w[:1].upper() + w[1:] if w else ""


# ----------------------------------------------------------------------------- имена
def nominative(gen):
    """имя в род. падеже («Никиты», «Степана», «Алексѣя», «Афонасья») -> именительный, если есть в словаре"""
    N = names()["M"]
    k = nk(gen)
    for v in (k[:-1], k[:-1] + "й", k[:-2] + "ий", k[:-1] + "а", k[:-1] + "я", k[:-2] + "ия", k):
        if len(v) >= 2 and v in N:
            return v
    return None


def patronymic(father, female):
    """отчество от имени отца в народной форме сказок: Никита -> Никитин, Гаврила -> Гаврилов,
    Василий -> Васильев, Алексей -> Алексеев, Пимен -> Пименов"""
    f = nk(father)
    if not f:
        return ""
    if f.endswith("ий"):
        base = f[:-2] + "ьев"
    elif f.endswith(("ей", "ай", "уй")):
        base = f[:-1] + "ев"
    elif f.endswith("ь"):
        base = f[:-1] + "ев"
    elif f.endswith(("ла", "ло")):          # Гаврила, Михайла, Данила
        base = f[:-1] + "ов"
    elif f.endswith(("а", "я")):            # Никита, Илья, Фома, Кузьма
        base = f[:-1] + "ин"
    else:
        base = f + "ов"
    base = base + ("а" if female else "")
    return base[:1].upper() + base[1:]


REL_NORM = {"сн": "сын", "снъ": "сын", "сына": "сын", "сыновья": "сын", "сыновя": "сын", "сынов": "сын",
            "дщ": "дочь", "дщерь": "дочь", "дочери": "дочь", "дочерь": "дочь", "дочи": "дочь",
            "братья": "брат", "брата": "брат", "сестры": "сестра", "внуки": "внук", "внучки": "внучка",
            "племянники": "племянник", "пасынки": "пасынок", "падчерицы": "падчерица",
            "приемышь": "приёмыш", "приемышъ": "приёмыш", "приемыш": "приёмыш", "салдатка": "солдатка",
            "вдов": "вдовец"}
CHILD = {"сын", "дочь", "дети", "детей"}
ORDINAL = {"второй", "вторый", "третий", "третии", "четвертый", "пятый", "шестой", "седьмой", "меньшой",
           "большой", "старший", "младший", "родной", "двоюродной", "двоюродный", "сводной", "сводный",
           "холостой", "женатый"}


def rel_word(tok, female):
    k = REL_NORM.get(nk(tok), nk(tok))
    if k in ("дети", "детей"):
        k = "дочь" if female else "сын"
    return k


# ----------------------------------------------------------------------------- значения ячеек
def parse_age(text):
    """-> (как показать, число|None). «55.» -> 55; «3 1/2» -> 3½; «мцу.» -> мес.; «новорожденные» -> новорожд."""
    t = modern(text or "").lower().strip(" .,:;")
    if not t:
        return "", None
    if "новорож" in t or "новрож" in t:
        return "новорожд.", None
    if re.search(r"\bм[еє]?с|мц|недел", t):
        return "мес.", 0
    if re.fullmatch(r"(1/2|½)", t):
        return "½", 0.5
    m = re.match(r"^(\d{1,3})?\s*(1/2|½)?$", t)
    if m and (m.group(1) or m.group(2)):
        n = int(m.group(1) or 0)
        if m.group(2):
            return (f"{n}½" if n else "½"), n + 0.5
        return n, n
    m = re.search(r"\d{1,3}", t)
    if m:
        return f"{int(m.group())}?", int(m.group())
    return f"{t}?", None


REASONS = [("умер", r"умер|умр|помер"), ("рекруты", r"рекрут"), ("переведён", r"перевед|переселен|переход"),
           ("бежал", r"беж|бег"), ("ссылка/поселение", r"сослан|ссыл|поселени"),
           ("отпущен", r"отпущ|уволен|вольн"), ("продан/куплен", r"продан|куплен|крепост")]


def parse_departure(text):
    """-> (примечание в современном виде, причина, год|None)"""
    t = modern(text or "").strip(" .;,")
    if not t:
        return "", "", None
    t = re.sub(r"\bв\s+(\d{4})", r"\1", t)
    t = re.sub(r"\s+", " ", t).strip(" .")
    years = re.findall(r"1[6-9]\d\d", t)
    reason = next((r for r, rx in REASONS if re.search(rx, t.lower())), "иное")
    if len(years) > 1 and len(set(years)) > 1 and reason in ("умер", "рекруты") and "или" not in t:
        t = f"{reason if reason == 'умер' else 'отдан в рекруты'} {' или '.join(years)}"
    return t, reason, (int(years[0]) if years else None)


# ----------------------------------------------------------------------------- селение
PLACE_RX = re.compile(r"\b(села|сельца|деревни|деревень|слободы|погоста|села|сёла|выселка|выселков|хутора)\s+"
                      r"([А-ЯЁ][\w-]+)", re.I)
KIND = {"села": "с.", "сельца": "сц.", "деревни": "д.", "деревень": "д.", "слободы": "сл.", "погоста": "пог.",
        "выселка": "выс.", "выселков": "выс.", "хутора": "х."}


def place_nom(kind, name):
    n = name
    if kind in ("д.", "сл."):        # Черепковой -> Черепкова, Рублевой -> Рублева
        n = re.sub(r"ой$", "а", n)
        n = re.sub(r"ей$", "я", n)
    elif kind in ("с.", "сц."):      # Троицкаго / Троицкого -> Троицкое
        n = re.sub(r"(аго|ого)$", "ое", n)
        n = re.sub(r"(яго|его)$", "ее", n)
    return f"{kind} {n}"


def detect_place(titles):
    """-> (селение, сословие|None, как в шапке) по текстам шапки страницы"""
    text = modern(" ".join(titles or []))
    m = PLACE_RX.search(text)
    estate = "дворовые" if re.search(r"дворов", text, re.I) else None
    if not m:
        return None, estate, ""
    kind = KIND.get(m.group(1).lower(), m.group(1).lower())
    return place_nom(kind, m.group(2)), estate, m.group(0)


def parse_places(spec):
    """«130-153=с. Троицкое;154-156=д. Екатериненская» -> [(130,153,'с. Троицкое'),…]"""
    out = []
    for part in (spec or "").split(";"):
        if "=" in part:
            rng, name = part.split("=", 1)
            a, _, b = rng.strip().partition("-")
            out.append((int(a), int(b or a), name.strip()))
    return out


def detect_year(titles):
    m = re.search(r"\b(1[78]\d\d)\s*год", modern(" ".join(titles or [])))
    return int(m.group(1)) if m else None


# ----------------------------------------------------------------------------- разбор сегмента
def parse_segment(seg, female):
    """«Никиты Григорьева второй сынъ Гаврила» -> dict(ref=[Никиты, Григорьева], rel='сын', ord=['второй'],
    name='Гаврила', rest=[...], pron=bool)"""
    toks = [t for t in seg.split() if re.search(r"\w", t)]
    keys = [nk(t) for t in toks]
    ri = next((i for i, k in enumerate(keys) if k in RELATION or k in REL_NORM), None)
    res = {"ref": [], "rel": "", "ord": [], "name": "", "rest": [], "pron": False, "raw": seg}
    if ri is None:
        if toks:
            res["name"], res["rest"] = toks[0], toks[1:]
        return res
    pre = toks[:ri]
    res["pron"] = any(nk(t) in PRON for t in pre + toks[ri + 1:ri + 3])
    res["ref"] = [t for t in pre if nk(t) not in PRON and nk(t) not in ORDINAL]
    res["ord"] = [modern(t).lower().strip(":.,") for t in pre if nk(t) in ORDINAL]
    res["rel"] = rel_word(toks[ri], female)
    j = ri + 1
    while j < len(toks) and (nk(toks[j]) in PRON or nk(toks[j]) in ORDINAL or nk(toks[j]) in RELATION):
        if nk(toks[j]) in ORDINAL:
            res["ord"].append(modern(toks[j]).lower().strip(":.,"))
        j += 1
    if j < len(toks):
        res["name"], res["rest"] = toks[j], toks[j + 1:]
    return res


# ----------------------------------------------------------------------------- основной проход
def default_args():
    return SimpleNamespace(split_x=None, fam_tol=0.025, near_tol=0.01, col_tol=0.05, tall=1.9, row_overlap=0.5)


def load_reports(inputs):
    """готовые pages.json от extract_rows.py или страницы __NEXT_DATA__"""
    reps = []
    for f in ER.collect_inputs(inputs):
        if f.suffix == ".json":
            try:
                d = json.loads(f.read_text(encoding="utf-8"))
            except Exception:
                d = None
            if isinstance(d, list) and d and isinstance(d[0], dict) and "blocks" in d[0]:
                reps += d
                continue
        try:
            rep = ER.process(ER.load_node(f), default_args())
        except Exception as e:
            rep = {"page": None, "flags": [f"error:{type(e).__name__}:{e}"], "blocks": [], "titles": []}
        rep["source"] = str(f)
        reps.append(rep)
    reps.sort(key=lambda r: (r.get("page") is None, r.get("page") or 0))
    return reps


def cell(row, role):
    c = row["cells"].get(role) or {}
    return c.get("text", "")


def build(reps, a):
    places = parse_places(a.places)
    year = a.year
    people, blocks_out, pages_out = [], [], []
    state = {"place": None, "estate": "крестьяне", "fam": {"M": None, "F": None}}
    fam_members = {}           # (селение, семья) -> [person]
    for rep in reps:
        page = rep.get("page")
        titles = rep.get("titles") or []
        year = year or detect_year(titles)
        place, estate, place_raw = detect_place(titles)
        if place and place != state["place"]:
            state.update(place=place, estate=estate or "крестьяне", fam={"M": None, "F": None})
        elif estate:
            state["estate"] = estate
        for lo, hi, nm in places:
            if page is not None and lo <= page <= hi:
                if nm != state["place"]:
                    state["fam"] = {"M": None, "F": None}
                state["place"] = nm
                state["estate"] = "дворовые" if "дворов" in nm else state["estate"]
        settle = state["place"] or "не определено"
        pflags = list(rep.get("flags") or [])
        tc = rep.get("totals_check") or {}
        pages_out.append({"page": page, "place": settle, "estate": state["estate"], "place_raw": place_raw,
                          "titles": " | ".join(titles), "flags": pflags, "totals": tc,
                          "url": a.url.format(page=page) if a.url and page is not None else ""})
        for bl in rep.get("blocks") or []:
            sex = bl["sex"]
            female = sex == "F"
            fno = bl.get("family_no")
            bflags = list(bl.get("flags") or [])
            if fno is None:
                fno = state["fam"][sex]
                if fno is not None:
                    bflags.append("семья с предыдущего скана")
            else:
                state["fam"][sex] = fno
            fkey = (settle, fno)
            members = fam_members.setdefault(fkey, [])
            rows = bl["rows"]
            segs = split_block(bl.get("names_text", ""), [sex] * len(rows))
            found = sum(1 for s, _ in segs if s)
            blocks_out.append({"page": page, "place": settle, "family": fno, "sex": "м" if not female else "ж",
                               "text": bl.get("names_text", ""), "rows": len(rows), "people": found,
                               "flags": bflags})
            last = {"rel": "", "ref": None, "ref_words": "", "ord": []}
            for r, (seg, sflag) in zip(rows, segs):
                p = person(seg, sflag, r, bl, bflags, page, settle, state["estate"], fno, female, members, last, year)
                p["url"] = pages_out[-1]["url"]
                members.append(p)
                people.append(p)
    return people, blocks_out, pages_out, year


def find_ref(ref_words, members):
    """ищет упомянутого человека среди уже названных в семье: «Никиты Григорьева» -> Никита Григорьев"""
    if not ref_words:
        return None
    nm = nominative(ref_words[0])
    if not nm:
        return None
    pat = nk(ref_words[1])[:5] if len(ref_words) > 1 else ""
    for p in reversed(members):
        if p["sex"] == "м" and nk(p["first"]) == nm:
            if not pat or nk(p["patr"]).startswith(pat[:4]) or not p["patr"]:
                return p
    for p in reversed(members):
        if p["sex"] == "м" and nk(p["first"]) == nm:
            return p
    return None


def person(seg, sflag, r, bl, bflags, page, settle, estate, fno, female, members, last, year):
    flags = [f for f in (bflags + r.get("flags", [])) if f]
    if sflag:
        flags.append(sflag)
    ps = parse_segment(seg, female)
    first = word(ps["name"])
    rest = [word(t) for t in ps["rest"] if word(t)]
    qual = [w.lower() for w in rest if nk(w) in ORDINAL]
    rest = [w for w in rest if nk(w) not in ORDINAL]
    rel, ref, ref_words, ords = ps["rel"], None, "", ps["ord"]
    if rel:
        if ps["ref"]:
            ref = find_ref(ps["ref"], members)
            ref_words = " ".join(word(t) for t in ps["ref"])
        elif ps["pron"] or not ps["ref"]:
            ref, ref_words = last["ref"], last["ref_words"]
        last.update(rel=rel, ref=ref, ref_words=ref_words, ord=ords)
    elif seg and last["rel"]:                         # «Степана сыновья: федоръ Николай …» — продолжение списка
        rel, ref, ref_words, ords = last["rel"], last["ref"], last["ref_words"], []
    head = None
    if not rel and seg:
        if not female and not any(m["rel"] == "глава" for m in members):
            rel = "глава"
        elif female and re.search(r"вдов", modern(seg).lower()):
            rel = "вдова, глава" if not members else "вдова"
        else:
            rel = "?"
    # отчество: своё из текста, иначе от отца (сын/дочь), брат — как у брата
    patr = rest[0] if rest else ""
    derived = False
    if not patr and rel in ("сын", "дочь") and (ref or ref_words):
        father = ref["first"] if ref else (nominative(ps["ref"][0]) if ps["ref"] else "")
        if father:
            patr, derived = patronymic(father, female), True
    if not patr and rel in ("брат", "сестра") and ref and ref.get("patr"):
        base = ref["patr"][:-1] if ref["sex"] == "ж" and ref["patr"].endswith("а") else ref["patr"]
        patr = base + ("а" if female else "")
        derived = True
    surname = " ".join(rest[1:]) if len(rest) > 1 else ""
    fio = " ".join(x for x in [first] + qual + [patr, surname] if x)
    if rel and rel not in ("глава", "?", "вдова", "вдова, глава"):
        who = ref_words or (f"({ref['fio']})" if ref else "")
        rel_text = " ".join(x for x in ords + [rel, who] if x)
    else:
        rel_text = rel
    if derived:
        flags.append("отчество по имени отца")
    # возраст и выбытие
    prev_d, prev_n = parse_age(cell(r, "age_prev"))
    now_d, now_n = parse_age(cell(r, "age_now"))
    doubt = [f for f in r.get("flags", []) if f.startswith(DOUBT_FLAGS)]
    if doubt:
        if prev_d != "" and "?" not in str(prev_d) and any("age_prev" in f or f in ("y_near",) for f in doubt):
            prev_d = f"{prev_d}?"
        if now_d != "" and "?" not in str(now_d) and any("age_now" in f or f in ("y_near",) for f in doubt):
            now_d = f"{now_d}?"
    note, reason, dep_year = parse_departure(cell(r, "departed_when"))
    absent = modern(cell(r, "absent_since")).strip(" .")
    status = "выбыл" if note else ("в отлучке" if absent else "налицо")
    if absent:
        note = (note + "; " if note else "") + f"в отлучке: {absent}"
    prev_year = REVISIONS.get(max((k for k, v in REVISIONS.items() if year and v < year), default=0))
    by = None
    if isinstance(now_n, (int, float)) and year:
        by = int(round(year - now_n))
    elif isinstance(prev_n, (int, float)) and prev_year:
        by = int(round(prev_year - prev_n))
    elif prev_d == "новорожд." and prev_year:
        by = f"после {prev_year}"
    pid = f"{page}-{'ж' if female else 'м'}{bl['block_no']}-{r['row_no']}"
    return {"id": pid, "page": page, "place": settle, "estate": estate, "family": fno,
            "sex": "ж" if female else "м", "first": first, "patr": patr, "surname": surname, "fio": fio,
            "rel": rel, "relation": rel_text, "ref_id": ref["id"] if ref else "",
            "age_prev": prev_d, "age_now": now_d, "birth": by, "status": status, "note": note,
            "reason": reason, "dep_year": dep_year, "raw": seg, "block_text": bl.get("names_text", ""),
            "flags": flags}


# ----------------------------------------------------------------------------- фамилия семьи
def family_surnames(people):
    heads = {}
    for p in people:
        if p["rel"] in ("глава", "вдова, глава") and p["surname"]:
            heads.setdefault((p["place"], p["family"]), p["surname"])
    for p in people:
        p["fam_surname"] = heads.get((p["place"], p["family"]), "")


# ----------------------------------------------------------------------------- Excel
def write_xlsx(path, people, blocks, pages, year, a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    prev_year = REVISIONS.get(max((k for k, v in REVISIONS.items() if year and v < year), default=0))
    prev_no = max((k for k, v in REVISIONS.items() if year and v < year), default=None)
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    warn = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, header, rows, widths, flag_col=None):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = bold, head_fill, wrap
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
        if flag_col:
            for r in ws.iter_rows(min_row=2):
                if r[flag_col - 1].value:
                    for c in r:
                        c.fill = warn

    ws = wb.active
    ws.title = "Люди"
    hdr = ["Скан", "Селение", "Сословие", "Семья №", "Фамилия семьи", "Пол", "ID", "ФИО (совр. орфография)",
           "Отношение", "ID отца / мужа",
           f"Лет по {prev_no}-й рев. ({prev_year})" if prev_year else "Лет по прошлой ревизии",
           f"Лет в {year}" if year else "Лет нынче", "Год рождения ≈", "Статус", "Выбытие / примечание",
           "Как в тексте", "Проверить", "Ссылка на скан"]
    rows = []
    for p in people:
        rows.append([p["page"], p["place"], p["estate"], p["family"], p["fam_surname"], p["sex"], p["id"],
                     p["fio"] or "(имя не найдено)", p["relation"], p["ref_id"], p["age_prev"] if p["age_prev"] != "" else None,
                     p["age_now"] if p["age_now"] != "" else None, p["birth"], p["status"], p["note"] or None,
                     p["raw"], "; ".join(human(f) for f in p["flags"] if human(f)) or None, p["url"] or None])
    sheet(ws, hdr, rows, [6, 18, 11, 7, 13, 5, 12, 30, 28, 12, 11, 9, 10, 10, 30, 40, 30, 30], flag_col=17)
    for r in ws.iter_rows(min_row=2):
        if r[17].value:
            r[17].hyperlink = r[17].value

    # Семьи
    fams = {}
    for p in people:
        f = fams.setdefault((p["place"], p["family"]), {"pages": set(), "head": "", "m": 0, "f": 0, "mn": 0, "fn": 0,
                                                       "sur": p["fam_surname"]})
        f["pages"].add(p["page"])
        if p["rel"] in ("глава", "вдова, глава") and not f["head"]:
            f["head"] = p["fio"]
        f["m" if p["sex"] == "м" else "f"] += 1
        if p["status"] != "выбыл":
            f["mn" if p["sex"] == "м" else "fn"] += 1
    ws = wb.create_sheet("Семьи")
    sheet(ws, ["Селение", "Семья №", "Сканы", "Глава", "Фамилия", "Мужчин в списке", "Женщин в списке",
               "Мужчин налицо", "Женщин налицо"],
          [[k[0], k[1], ", ".join(str(x) for x in sorted(v["pages"], key=lambda z: z or 0)), v["head"], v["sur"],
            v["m"], v["f"], v["mn"], v["fn"]] for k, v in fams.items()],
          [18, 8, 12, 30, 14, 10, 10, 10, 10])

    ws = wb.create_sheet("Блоки текста")
    sheet(ws, ["Скан", "Селение", "Семья №", "Пол", "Текст блока (как распознан)", "Строк в таблице",
               "Людей найдено в тексте", "Проверить"],
          [[b["page"], b["place"], b["family"], b["sex"], b["text"], b["rows"], b["people"],
            "; ".join(human(f) for f in b["flags"] if human(f)) or
            ("людей в тексте меньше, чем строк" if b["people"] < b["rows"] else None)] for b in blocks],
          [6, 18, 7, 5, 80, 9, 9, 30], flag_col=8)

    ws = wb.create_sheet("Страницы")
    prow = []
    for pg in pages:
        tc = pg["totals"] or {}
        def tot(s, k):
            return (tc.get(s) or {}).get(k)
        ok = all(v.get("ok") for v in tc.values()) if tc else None
        prow.append([pg["page"], pg["place"], pg["estate"], tot("M", "expected"), tot("M", "rows_present"),
                     tot("F", "expected"), tot("F", "rows_present"),
                     "да" if ok else ("нет" if ok is False else "итогов нет"),
                     "; ".join(human(f) for f in pg["flags"] if human(f)) or None, pg["titles"], pg["url"] or None])
    sheet(ws, ["Скан", "Селение", "Сословие", "Итог М (на скане)", "М налицо по строкам", "Итог Ж (на скане)",
               "Ж налицо по строкам", "Итоги сошлись", "Проверить", "Шапка страницы", "Ссылка"],
          prow, [6, 18, 11, 9, 9, 9, 9, 10, 30, 60, 30], flag_col=9)

    ws = wb.create_sheet("Сводка")
    pl = sorted({p["place"] for p in people}, key=[p["place"] for p in people].index)
    n = len(people) + 1
    ws.append(["Селение", "Мужчин в списке", "Женщин в списке", "Всего", "Мужчин налицо", "Женщин налицо",
               "Выбыло"])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    for i, name in enumerate(pl, 2):
        rng = f"Люди!$B$2:$B${n}"
        sx = f"Люди!$F$2:$F${n}"
        st = f"Люди!$N$2:$N${n}"
        ws.append([name, f'=COUNTIFS({rng},A{i},{sx},"м")', f'=COUNTIFS({rng},A{i},{sx},"ж")', f"=B{i}+C{i}",
                   f'=COUNTIFS({rng},A{i},{sx},"м",{st},"<>выбыл")', f'=COUNTIFS({rng},A{i},{sx},"ж",{st},"<>выбыл")',
                   f'=COUNTIFS({rng},A{i},{st},"выбыл")'])
    k = len(pl) + 2
    ws.append(["Итого"] + [f"=SUM({c}2:{c}{k - 1})" for c in "BCDEFG"])
    for c in ws[k]:
        c.font = bold
    for col, w in zip("ABCDEFG", [20, 12, 12, 9, 12, 12, 9]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Об источнике")
    lines = []
    if a.source:
        lines.append(f"Источник: {a.source}")
    if pages and pages[0]["titles"]:
        lines.append(f"Шапка первого скана: {pages[0]['titles']}")
    places_line = []
    for name in pl:
        pgs = sorted({p["page"] for p in people if p["place"] == name and p["page"] is not None})
        if pgs:
            places_line.append(f"{name} ({pgs[0]}–{pgs[-1]})" if len(pgs) > 1 else f"{name} ({pgs[0]})")
    if places_line:
        lines.append("Селения: " + ", ".join(places_line) + ".")
    lines += [
        "Как собрано: из готового машинного распознавания Яндекс Архива (данные страницы, без повторного "
        "распознавания сканов). Текст привязан к столбцам таблицы по координатам блоков, список людей семьи "
        "разделён по именам и сопоставлен со строками таблицы сверху вниз.",
        "Орфография имён приведена к современной (без ъ, ѣ, і, ѳ); отчества и фамилии — как в документе. "
        "Если отчества в тексте нет, оно образовано от имени отца (отмечено в «Проверить»).",
        "«Отношение» — как в сказке: к кому относится человек (глава семьи, сын/жена/дочь названного лица). "
        "«ID отца / мужа» связывает людей между собой для родословного древа.",
        f"«новорожд.» в столбце прошлой ревизии — родился после {prev_year or 'прошлой ревизии'}. "
        "«½» — полгода, «мес.» — несколько месяцев. «?» — сомнительное значение (ячейка у границы строки "
        "или несколько значений в одном месте); пустая ячейка — цифра не распознана. Год рождения ≈ = год ревизии "
        "минус возраст.",
        "Статус: «выбыл» — есть отметка о выбытии (умер, рекруты, переведён), «в отлучке» — отметка в женском "
        "столбце отлучки, иначе «налицо». Номер семьи — номер в пределах селения.",
        "Строки, выделенные жёлтым, стоит сверить со сканом (ссылка в последнем столбце).",
    ]
    ws.column_dimensions["A"].width = 140
    for ln in lines:
        ws.append([ln])
        ws.cell(ws.max_row, 1).alignment = wrap
    wb.save(path)


HUMAN = {
    "семья с предыдущего скана": "номер семьи взят с предыдущего скана",
    "имя в тексте блока не найдено": "имя в тексте не найдено",
    "пол по имени не совпал со столбцом": "пол по имени не совпал со столбцом",
    "отчество по имени отца": "отчество образовано от имени отца",
    "nested_block": "", "family_missing": "номер семьи не найден",
    "family_changes_inside_block": "номер семьи меняется внутри блока",
    "y_near": "значение у границы строки", "weak_col": "значение у границы столбца",
}


def human(f):
    if f in HUMAN:
        return HUMAN[f]
    if f.startswith("span:"):
        return ""
    if f.startswith("multi_cell:") or f.startswith("span_conflict:"):
        return "несколько значений в одном месте: " + f.split(":", 1)[1]
    if f.startswith("totals_mismatch"):
        return "строки не сошлись с итогом страницы"
    if f.startswith("count_mismatch"):
        return "данные страницы неполные — страница пропущена"
    if f.startswith("no_main_header"):
        return "форма не распознана (нет столбца «Крестьяне»)"
    if f.startswith("orphans"):
        return "есть неотнесённые блоки текста"
    return f


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", help="страницы (.json/.html), папки или pages.json от extract_rows.py")
    ap.add_argument("--out-xlsx", required=True, help="куда сохранить Excel")
    ap.add_argument("--out-json", default=None, help="люди в JSON (по умолчанию рядом с .xlsx)")
    ap.add_argument("--year", type=int, default=None, help="год сказки (по умолчанию из шапки)")
    ap.add_argument("--places", default="", help="селения по сканам, если шапка не распознана: "
                                                "«130-153=с. Троицкое;154-156=д. Екатериненская»")
    ap.add_argument("--url", default="", help="шаблон ссылки на скан, например "
                                             "https://ya.ru/archive/catalog/<id документа>/{page}")
    ap.add_argument("--source", default="", help="строка «Источник» (архив, фонд, опись, дело)")
    a = ap.parse_args()
    reps = load_reports(a.inputs)
    people, blocks, pages, year = build(reps, a)
    family_surnames(people)
    write_xlsx(a.out_xlsx, people, blocks, pages, year, a)
    oj = a.out_json or str(Path(a.out_xlsx).with_suffix(".people.json"))
    Path(oj).write_text(json.dumps(people, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    nf = sum(1 for p in people if not p["first"])
    print(f"сканов {len(pages)}, людей {len(people)} (без имени {nf}), год {year} -> {a.out_xlsx}, {oj}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
