"""
yaarchive: запись итога — Итог_люди.xlsx, Итог_древо.ged (GEDCOM 5.5.1), Итог_люди.json, Отчёт.txt
и снимок значений (по нему навык «Исправление ошибок» находит ваши ручные правки в Excel).
"""
import datetime
import json
import re
import time
from collections import Counter, OrderedDict
from pathlib import Path

XLSX, GED, JSONF, REPORT = "Итог_люди.xlsx", "Итог_древо.ged", "Итог_люди.json", "Отчёт.txt"
SNAP, PROPS = ".снимок.json", ".предложения.json"

# столбцы листа «Люди»: (заголовок, поле, ширина, можно править)
COLS = [
    ("Год", "year", 7, False), ("Документ", "doc_kind", 14, False), ("Населённый пункт", "settlement", 9, True),
    ("Двор / семья №", "yard", 8, True), ("Номер", "no", 7, True), ("Человек (как распознано)", "raw", 42, True),
    ("Имя", "first", 12, True), ("Отчество", "patr", 14, True), ("Фамилия", "surname", 12, True), ("Пол", "sex", 5, True),
    ("Возраст", "age", 8, True), ("Возраст по прошлой ревизии", "age_prev", 10, True), ("Год рождения", "birth", 9, True),
    ("Родство", "relation", 24, True), ("Статус", "status", 10, False), ("Выбытие / примечание", "note", 22, False),
    ("Проверить", "flags_text", 40, False), ("Исправлено", "fixed_text", 30, False), ("ID", "id", 22, False),
    ("ID отца / мужа", "ref_id", 22, True), ("Ссылка на страницу", "url", 18, False),
    ("Локальная Ссылка на страницу", "scan_link", 30, False), ("Архив хранения документа", "cga", 40, False), ("Ключ", "key", 10, False),
]
# цвета (светлые, кроме красного) — как в расширении; справочник («цвета») может заменить
COLORS = {"мужчины": "DDEBF7", "женщины": "F6E1F5", "ошибка": "FFC7CE", "НП_не_определён": "EDEDED",
          "НП": {"село Троицкое": "E2EFDA", "деревня Острогина": "FFF2CC", "деревня Черепкова": "E4DFEC",
                 "деревня Рублева": "D0ECE7", "деревня Луки": "FBE5D6", "деревня Мякинина": "EBF1DE",
                 "деревня Екатериненская": "F2E8D5"},
          "НП_прочие": ["E8F6EF", "F3EEDB", "EDE7F6", "E0F2F1", "FFF8E1", "F1F8E9", "E3F2FD", "FBE9E7"],
          "дворы": ["FFF9E5", "EAF4DD", "F1ECF7", "E6F3F6", "FBF0E4", "EEF0F3"],
          "слитые": ["F2F2F2", "E9EFF7", "EEF5E9", "F7F0E8"]}


def settlement_color(name):
    """цвет закреплён за названием НП: из справочника, иначе — по названию (одинаковый при каждом запуске)"""
    n = str(name or "")
    if not n or n.startswith("не определён"):
        return COLORS["НП_не_определён"]
    base = re.sub(r"^чужой НП:\s*", "", n)
    if base in COLORS["НП"]:
        return COLORS["НП"][base]
    h = 0
    for ch in base:
        h = (h * 31 + ord(ch)) & 0xFFFFFFFF
    return COLORS["НП_прочие"][h % len(COLORS["НП_прочие"])]


def error_fields(flags):
    """в какие ячейки указывают замечания «Проверить» (они заливаются светло-красным)"""
    out = {"flags_text"} if flags else set()
    for t in flags:
        t = t.lower()
        if "двор" in t:
            out.add("yard")
        elif "номер" in t:
            out.add("no")
        if "без номера" in t:              # «возраст без номера человека» — ошибка в номере, не в возрасте
            continue
        if "возраст" in t or "год рождения" in t:
            out.add("age")
        if "год рождения" in t:
            out.add("birth")
        if any(w in t for w in ("имя", "пол по имени", "текст", "блок", "привязк")):
            out.add("raw")
        if "родств" in t or "отец" in t:
            out.add("relation")
    return out
EDITABLE = {c[1]: c[0] for c in COLS if c[3]}


def val(p, f):
    v = p.get(f, "")
    if f == "no" and str(v).isdigit():
        return int(v)
    if f in ("age", "age_prev") and re.fullmatch(r"\d+", str(v or "")):
        return int(v)
    return "" if v is None else v


def write_all(ctx, od, KIND_RU, rel_link, version):
    people = ctx.people
    page_of = {}
    for pg in ctx.pages:
        page_of[(pg.doc["id"], pg.page)] = pg
    for p in people:
        pg = page_of.get((p["doc"], p["page"]))
        p["doc_kind"] = KIND_RU[p["kind"]]
        p["flags_text"] = "; ".join(p["flags"])
        p["fixed_text"] = "; ".join(p["fixed"])
        p["cga"] = pg.ci["cga"] if pg and pg.ci else ""
        p["cga_url"] = pg.ci["cga_url"] if pg and pg.ci else ""
        p["scan_link"] = rel_link(pg.scan_path, od) if pg and pg.scan_path else ""
    res = {}
    res["xlsx"], res["locked"] = write_excel(ctx, od, KIND_RU, version)
    res["fams"] = write_gedcom(od / GED, ctx, KIND_RU, version)
    write_json(od / JSONF, ctx, KIND_RU, version)
    write_report(od / REPORT, ctx, KIND_RU, version, res)
    snap = {"written": time.time(), "xlsx": str(res["xlsx"]), "version": version, "root": str(ctx.roots[0]),
            "values": {p["key"]: {f: str(val(p, f)) for f in EDITABLE} for p in people}}
    (od / SNAP).write_text(json.dumps(snap, ensure_ascii=False), encoding="utf-8")
    props = {pr["id"]: pr for pr in ctx.proposals}
    (od / PROPS).write_text(json.dumps(props, ensure_ascii=False, default=str), encoding="utf-8")
    return res


# =============================================================================== Excel
def write_excel(ctx, od, KIND_RU, version):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.merge import MergedCellRange
    fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
    F_HEAD, F_NP, F_YARD, F_CALC = fill("D9D9D9"), fill("E2EFDA"), fill("FFF2CC"), fill("FCE4D6")
    F_CHECK, F_FIX, F_ASK = fill(COLORS["ошибка"]), fill("DDEBF7"), fill("FFFF99")
    thin = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    BOLD, LINK = Font(bold=True), Font(color="0563C1", underline="single")
    WRAP = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws = wb.active
    ws.title = "Люди"
    ws.append([c[0] for c in COLS])
    for i, c in enumerate(COLS, 1):
        cell = ws.cell(1, i)
        cell.font, cell.fill, cell.alignment = BOLD, F_HEAD, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = c[2]
    ws.column_dimensions[ws.cell(1, len(COLS)).column_letter].hidden = True
    ci = {c[1]: i for i, c in enumerate(COLS, 1)}
    people = ctx.people
    for p in people:
        ws.append([val(p, c[1]) for c in COLS])
    n = len(people)
    # постранично: одна слитая ячейка «Год», «Ссылка на страницу», «Локальная Ссылка…», «Архив хранения документа»
    page_key = lambda p: (p["doc"], p["page"])  # noqa: E731
    page_run, yard_run, k1, k2 = [], [], -1, -1
    for i, p in enumerate(people):
        if i == 0 or page_key(p) != page_key(people[i - 1]):
            k1 += 1
        page_run.append(k1)
        yk = (p["doc"], p["settlement"], p["yard"], p["yard_calc"]) if p["yard"] != "" else ("·", i)
        if i == 0 or yk != ((people[i - 1]["doc"], people[i - 1]["settlement"], people[i - 1]["yard"], people[i - 1]["yard_calc"])
                            if people[i - 1]["yard"] != "" else ("·", i - 1)):
            k2 += 1
        yard_run.append(k2)
    fill_cache = {}

    def F(rgb):
        if rgb not in fill_cache:
            fill_cache[rgb] = fill(rgb)
        return fill_cache[rgb]
    C_NP = Alignment(horizontal="center", vertical="center", text_rotation=90, wrap_text=True)
    C_MID = Alignment(horizontal="center", vertical="center")
    C_V = Alignment(vertical="center", wrap_text=True)
    ITAL = Font(bold=True, italic=True)
    page_cols = {ci["year"], ci["url"], ci["scan_link"], ci["cga"]}
    for r, p in enumerate(people, 2):
        errs = error_fields(p["flags"])
        sex_fill = F(COLORS["женщины"] if p["sex"] == "ж" else COLORS["мужчины"]) if p["sex"] in ("м", "ж") else None
        for f, c in ci.items():
            cell = ws.cell(r, c)
            if f == "settlement":
                cell.fill, cell.font, cell.border, cell.alignment = F(settlement_color(p["settlement"])), BOLD, BORDER, C_NP
            elif f == "yard":
                cell.fill = F(COLORS["ошибка"]) if "yard" in errs else F(COLORS["дворы"][yard_run[r - 2] % len(COLORS["дворы"])])
                cell.font, cell.border, cell.alignment = (ITAL if p["yard_calc"] else BOLD), BORDER, C_MID
            elif c in page_cols:
                cell.fill, cell.border, cell.alignment = F(COLORS["слитые"][page_run[r - 2] % len(COLORS["слитые"])]), BORDER, \
                    (C_MID if f == "year" else C_V)
            elif f == "key":
                continue
            else:
                cell.border = BORDER
                if f in errs:
                    cell.fill = F(COLORS["ошибка"])
                elif sex_fill is not None:
                    cell.fill = sex_fill
                cell.alignment = C_V if f in ("raw", "flags_text", "fixed_text", "note", "relation") else C_MID
        if p["url"]:
            c = ws.cell(r, ci["url"])
            c.value, c.hyperlink, c.font = f"стр. {p['page']}", p["url"], LINK
        if p["scan_link"]:
            c = ws.cell(r, ci["scan_link"])
            c.value, c.hyperlink, c.font = Path(p["scan_link"]).name, p["scan_link"], LINK

    def merge_runs(col, key):
        st = 0
        for i in range(1, n + 1):
            if i == n or key(people[i]) != key(people[st]):
                r1, r2 = st + 2, i + 1
                if r2 > r1:
                    for rr in range(r1 + 1, r2 + 1):
                        ws.cell(rr, col).value = None
                    mcr = MergedCellRange(ws, f"{ws.cell(r1, col).coordinate}:{ws.cell(r2, col).coordinate}")
                    ws.merged_cells.ranges.add(mcr)       # без квадратичной проверки openpyxl: диапазоны не пересекаются
                    ws._clean_merge_range(mcr)
                st = i
    if people:
        merge_runs(ci["settlement"], lambda p: (p["doc"], p["settlement"]))
        merge_runs(ci["yard"], lambda p: (p["doc"], p["settlement"], p["yard"], p["yard_calc"]) if p["yard"] != "" else id(p))
        for f in ("year", "url", "scan_link", "cga"):
            merge_runs(ci[f], page_key)
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLS) - 1).column_letter}{len(people) + 1}"

    # ---- Исправления
    wf = wb.create_sheet("Исправления")
    head = ["№", "ID предложения", "Что", "Где (год, стр., человек)", "Сейчас", "Предлагается", "Почему",
            "Применить? (да / нет)", "Ссылка на страницу"]
    wf.append(["Навык «Исправление ошибок». Впишите «да» или «нет» в жёлтый столбец и сохраните файл — скрипт сам "
               "применит решения и пересчитает итог. Ваши правки прямо в листе «Люди» тоже появятся здесь для подтверждения."])
    wf.append(head)
    for i, w in enumerate([5, 12, 22, 40, 26, 26, 50, 14, 16], 1):
        wf.column_dimensions[wf.cell(2, i).column_letter].width = w
        wf.cell(2, i).font, wf.cell(2, i).fill = BOLD, F_HEAD
    dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
    wf.add_data_validation(dv)
    for n, pr in enumerate(ctx.proposals, 1):
        wf.append([n, pr["id"], pr["type"], pr["where"], pr["now"], pr["new"], pr["reason"], "", ""])
        r = wf.max_row
        if pr.get("url"):
            c = wf.cell(r, 9)
            c.value, c.hyperlink, c.font = f"стр. {pr.get('page')}", pr["url"], LINK
        wf.cell(r, 8).fill = F_ASK
        dv.add(wf.cell(r, 8))
        for k in (4, 5, 6, 7):
            wf.cell(r, k).alignment = WRAP
    hist = ctx.corr.history()
    if hist:
        wf.append([])
        wf.append(["", "", "Применённые и отклонённые (впишите «отменить», чтобы снять решение)"])
        wf.cell(wf.max_row, 3).font = BOLD
        wf.append(["№", "ID", "Что", "Где", "Было", "Стало", "Когда / откуда", "Решение", ""])
        for i in range(1, 9):
            wf.cell(wf.max_row, i).font, wf.cell(wf.max_row, i).fill = BOLD, F_HEAD
        for n, h in enumerate(hist, 1):
            wf.append([n, h["id"], h["type"], h["where"], h["now"], h["new"], h["when"], h["decision"], ""])
            for k in (4, 5, 6):
                wf.cell(wf.max_row, k).alignment = WRAP
    wf.freeze_panes = "A3"

    # ---- Семьи
    wsf = wb.create_sheet("Семьи")
    wsf.append(["Год", "Документ", "Населённый пункт", "Двор / семья №", "Людей", "Мужчин", "Женщин", "Глава", "Состав"])
    fams = OrderedDict()
    for p in people:
        fams.setdefault((p["year"], p["doc_kind"], p["settlement"], p["yard"]), []).append(p)
    for (y, k, s, yd), mm in fams.items():
        head_p = next((m for m in mm if m["rel"] in ("глава", "вдова, глава")), mm[0])
        wsf.append([y, k, s, yd, len(mm), sum(m["sex"] == "м" for m in mm), sum(m["sex"] == "ж" for m in mm),
                    head_p["fio"] or head_p["raw"], "; ".join(f"{m['fio'] or m['raw'] or '?'} ({m['relation'] or '?'}, {m['age']})" for m in mm)])
    _head(wsf, BOLD, F_HEAD, [6, 18, 20, 8, 7, 7, 7, 26, 100])

    # ---- Страницы
    wsp = wb.create_sheet("Страницы")
    wsp.append(["Год", "Документ", "Тип определён", "Страница", "Людей", "Замечания", "Ссылка", "Скан на диске", "Пакет"])
    for pg in ctx.pages:
        wsp.append([pg.year, KIND_RU[pg.kind], pg.doc["kind_how"], pg.page, len(pg.people), "; ".join(pg.flags),
                    pg.url, pg.ci["scan_name"] if pg.scan_path else "нет", pg.pkg.name])
        if pg.url:
            c = wsp.cell(wsp.max_row, 7)
            c.hyperlink, c.font = pg.url, LINK
        if pg.flags:
            wsp.cell(wsp.max_row, 6).fill = F_CHECK
    _head(wsp, BOLD, F_HEAD, [6, 18, 26, 8, 7, 60, 30, 30, 20])

    # ---- Проверка
    wsc = wb.create_sheet("Проверка")
    wsc.append(["Год", "Страница", "Человек", "ID", "Что проверить", "Ссылка"])
    for p in people:
        if p["flags"]:
            wsc.append([p["year"], p["page"], p["fio"] or p["raw"] or "(без имени)", p["id"], "; ".join(p["flags"]), p["url"]])
    _head(wsc, BOLD, F_HEAD, [6, 8, 30, 22, 80, 30])

    # ---- Сводка
    wss = wb.create_sheet("Сводка")
    wss.append(["Год", "Документ", "Тип определён", "Дело", "Страниц", "Пропущено страниц", "Не загрузились",
                "Людей", "Мужчин", "Женщин", "НП", "Дворов / семей", "С годом рождения", "На проверку", "Сканов на диске"])
    for d in ctx.docs:
        pp = [p for p in people if p["doc"] == d["id"]]
        wss.append([d["year"], KIND_RU[d["kind"]], d["kind_how"], (d.get("ci") or {}).get("source", ""), len(d["pages"]),
                    _ranges(d["gaps"]), _ranges(sorted(f.get("page") for f in d["failed"] if f.get("page") is not None)),
                    len(pp), sum(p["sex"] == "м" for p in pp), sum(p["sex"] == "ж" for p in pp),
                    len({p["settlement"] for p in pp}), len({(p["settlement"], p["yard"]) for p in pp}),
                    sum(isinstance(p["birth"], int) for p in pp), sum(bool(p["flags"]) for p in pp),
                    sum(bool(pg.scan_path) for pg in d["pages"])])
    _head(wss, BOLD, F_HEAD, [6, 18, 26, 30, 8, 14, 12, 8, 8, 8, 6, 8, 10, 10, 10])

    # ---- Об обработке
    wa = wb.create_sheet("Об обработке")
    for line in about_lines(ctx, version):
        wa.append([line])
    wa.column_dimensions["A"].width = 130

    wb.move_sheet("Исправления", offset=-(len(wb.sheetnames) - 2))
    target = od / XLSX
    try:
        wb.save(target)
        return target, False
    except PermissionError:                      # файл открыт в Excel (Windows)
        alt = od / f"Итог_люди (новый {datetime.datetime.now():%H-%M}).xlsx"
        wb.save(alt)
        return alt, True


def _head(ws, bold, fill_, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(1, i)
        c.font, c.fill = bold, fill_
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _ranges(nums):
    out, nums = [], sorted(set(n for n in nums if n is not None))
    i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}" if i == j else f"{nums[i]}–{nums[j]}")
        i = j + 1
    return ", ".join(out)


def about_lines(ctx, version):
    return [
        f"yaarchive {version}, {datetime.datetime.now():%d.%m.%Y %H:%M}",
        f"Папка: {ctx.roots[0]}",
        f"Пакетов расширения обработано: {len(ctx.packages)}; ждут окончания записи: {len(ctx.waiting)}; "
        f"страниц, загруженных повторно (взяты из самого нового пакета): {ctx.dups}",
        "",
        "Порядок обработки: пакеты → страницы → документы (год + дело, по номеру страницы) → тип документа → разбор "
        "(исповедные: столбцы, номера по соседям, люди, НП и двор по порядку; ревизские: строки, люди, родство, выбытие) → "
        "НП по справочнику → принятые исправления значений → пустой двор по соседям → родство и отчество → принятые "
        "исправления родства → год рождения → сканы → проверка → навык «Исправление ошибок» → итог.",
        "",
        "Лист «Люди»: у всех людей одной страницы одна слитая ячейка «Год», «Ссылка на страницу», «Локальная Ссылка на "
        "страницу» и «Архив хранения документа» (текст). Цвета (светлые, кроме красного): у каждого НП свой цвет, закреплённый за "
        "названием; дворы и блоки страниц — чередуются; строки мужчин — голубые, женщин — розовые; ячейки с ошибками — "
        "светло-красные; двор курсивом — вычислен по соседним. «Не определён» НП получает название следующего определённого; "
        "повторный НП в документе перенесён сразу после первичного.",
        "Год рождения = год документа − возраст (±1 год). В GEDCOM каждая запись документа — отдельная персона: один и тот же "
        "человек в разные годы не объединён.",
        "Исправления: лист «Исправления» («да»/«нет», сохранить файл) или диалог в окне скрипта. Правки, сделанные прямо в листе "
        "«Люди», скрипт заметит и попросит подтвердить; принятые хранятся в «Итог/исправления.json» и применяются при каждом пересчёте.",
        "Справочники (имена, слова родства, населённые пункты и их порядок, годы ревизий, пороги проверок) — справочники.json; "
        "свою копию можно положить рядом со скриптом.",
    ]


# =============================================================================== GEDCOM
def ged(s):
    return re.sub(r"[\r\n]+", " ", str(s if s is not None else "")).strip()


def conc(line, limit=240):
    """строки длиннее предела GEDCOM 5.5.1 переносятся через CONC (не по пробелу — пробел не теряется)"""
    if len(line) <= limit:
        return line
    lvl = int(line.split(" ", 1)[0])
    out, rest = [line[:limit]], line[limit:]
    while rest:
        out.append(f"{lvl + 1} CONC {rest[:limit - 10]}")
        rest = rest[limit - 10:]
    return "\n".join(out)


def write_gedcom(path, ctx, KIND_RU, version):
    people = ctx.people
    ids = {p["id"]: f"@I{i}@" for i, p in enumerate(people, 1)}
    by_id = {p["id"]: p for p in people}
    fams, by_parents, famc = [], {}, {}

    def fam(h, w):
        if (h, w) not in by_parents:
            by_parents[(h, w)] = {"id": f"@F{len(fams) + 1}@", "HUSB": h, "WIFE": w, "CHIL": []}
            fams.append(by_parents[(h, w)])
        return by_parents[(h, w)]
    wife_of = {}
    for p in people:
        if p["rel"] == "жена" and p["ref_id"] in by_id and by_id[p["ref_id"]]["sex"] == "м":
            wife_of.setdefault(p["ref_id"], p["id"])
            fam(p["ref_id"], p["id"])
    for p in people:
        ref = by_id.get(p["ref_id"])
        if not ref:
            continue
        if p["rel"] in ("сын", "дочь"):
            f = fam(ref["id"], wife_of.get(ref["id"])) if ref["sex"] == "м" else fam(None, ref["id"])
            f["CHIL"].append(p["id"])
            famc[p["id"]] = f["id"]
        elif p["rel"] in ("брат", "сестра"):
            fid = famc.get(ref["id"])
            f = next((x for x in fams if x["id"] == fid), None) if fid else None
            if not f:
                f = {"id": f"@F{len(fams) + 1}@", "HUSB": None, "WIFE": None, "CHIL": [ref["id"]]}
                fams.append(f)
                famc[ref["id"]] = f["id"]
            f["CHIL"].append(p["id"])
            famc[p["id"]] = f["id"]
    fams_of = {}
    for f in fams:
        for role in ("HUSB", "WIFE"):
            if f[role]:
                fams_of.setdefault(f[role], []).append(f["id"])
    t = datetime.date.today()
    mon = "JAN FEB MAR APR MAY JUN JUL AUG SEP OCT NOV DEC".split()[t.month - 1]
    L = ["0 HEAD", "1 SOUR YAARCHIVE", f"2 VERS {version}", "2 NAME yaarchive (Яндекс Архив)", f"1 DATE {t.day} {mon} {t.year}",
         "1 GEDC", "2 VERS 5.5.1", "2 FORM LINEAGE-LINKED", "1 CHAR UTF-8", "1 LANG Russian",
         "1 NOTE Люди из исповедных ведомостей и ревизских сказок (Яндекс Архив). Каждая запись документа — отдельная "
         "персона: один и тот же человек в разные годы не объединён. Год рождения — по возрасту (±1 год)."]
    src = {}
    for k, d in enumerate(ctx.docs, 1):
        src[d["id"]] = f"@S{k}@"
        ci = d.get("ci") or {}
        L += [f"0 @S{k}@ SOUR", f"1 TITL {KIND_RU[d['kind']]} {d['year'] or ''} г."]
        if ci.get("source"):
            L.append(f"1 AUTH {ged(ci['source'])}")
        if ci.get("file_title"):
            L.append(f"1 PUBL {ged(ci['file_title'])}")
        if ci.get("cga_url"):
            L.append(f"1 NOTE {ci['cga_url']}")
    for p in people:
        given = " ".join(x for x in (p["first"], p["patr"]) if x) or "(имя не прочитано)"
        L += [f"0 {ids[p['id']]} INDI", f"1 NAME {ged(given)} /{ged(p['surname'])}/", f"2 GIVN {ged(given)}"]
        if p["surname"]:
            L.append(f"2 SURN {ged(p['surname'])}")
        L.append(f"1 SEX {'M' if p['sex'] == 'м' else 'F'}")
        b = p["birth"]
        if isinstance(b, int):
            L += ["1 BIRT", f"2 DATE ABT {b}", f"2 NOTE По возрасту {ged(p['age'] or p['age_prev'])} в {p['year']} г."]
        elif isinstance(b, str) and b.startswith("после "):
            L += ["1 BIRT", f"2 DATE AFT {b.split()[-1]}"]
        L += ["1 CENS"] + ([f"2 DATE {p['year']}"] if p["year"] else []) + [f"2 PLAC {ged(p['settlement'])}"]
        det = [f"двор/семья {p['yard']}" + (" (вычислен)" if p["yard_calc"] else "") if p["yard"] != "" else "",
               f"№ {p['no']}" if p["no"] not in ("", None) else "", f"возраст {p['age']}" if p["age"] not in ("", None) else "",
               p["relation"] if p["relation"] not in ("", "?") else "", p["note"]]
        L += [f"2 NOTE {ged('; '.join(x for x in det if x))}", f"2 SOUR {src[p['doc']]}", f"3 PAGE Скан {p['page']}",
              "3 DATA", f"4 TEXT {ged(p['raw'])}"]
        if p["dep_reason"] == "умер" and str(p["dep_year"]).isdigit():
            L += ["1 DEAT", f"2 DATE {p['dep_year']}", f"2 SOUR {src[p['doc']]}"]
        for fid in fams_of.get(p["id"], []):
            L.append(f"1 FAMS {fid}")
        if p["id"] in famc:
            L.append(f"1 FAMC {famc[p['id']]}")
        if p["url"]:
            L.append(f"1 NOTE Скан: {p['url']}")
        if p["flags"]:
            L.append(f"1 NOTE Проверить: {ged('; '.join(p['flags']))}")
    for f in fams:
        L.append(f"0 {f['id']} FAM")
        for role in ("HUSB", "WIFE"):
            if f[role]:
                L.append(f"1 {role} {ids[f[role]]}")
        for c in OrderedDict.fromkeys(f["CHIL"]):
            L.append(f"1 CHIL {ids[c]}")
    L.append("0 TRLR")
    Path(path).write_text("\n".join(conc(x) for x in L) + "\n", encoding="utf-8")
    return len(fams)


# =============================================================================== JSON и отчёт
def write_json(path, ctx, KIND_RU, version):
    skip = {"locked", "flags_text", "fixed_text"}
    data = {"version": version, "created": datetime.datetime.now().isoformat(timespec="seconds"),
            "documents": [{"id": d["id"], "year": d["year"], "kind": KIND_RU[d["kind"]], "source": (d.get("ci") or {}).get("source", ""),
                           "pages": [p.page for p in d["pages"]], "gaps": d["gaps"]} for d in ctx.docs],
            "people": [{k: v for k, v in p.items() if k not in skip} for p in ctx.people]}
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


def write_report(path, ctx, KIND_RU, version, res):
    people = ctx.people
    flags = Counter(f.split(":")[0] if f.startswith("лист") is False else f for p in people for f in p["flags"])
    L = [f"Отчёт yaarchive {version} — {datetime.datetime.now():%d.%m.%Y %H:%M}", f"Папка: {ctx.roots[0]}", "",
         f"Пакетов: {len(ctx.packages)} (ждут записи: {len(ctx.waiting)}), страниц: {len(ctx.pages)}, "
         f"повторно загруженных страниц: {ctx.dups}", f"Людей: {len(people)}, документов: {len(ctx.docs)}", ""]
    for d in ctx.docs:
        pp = [p for p in people if p["doc"] == d["id"]]
        L.append(f"• {d['year']} — {KIND_RU[d['kind']]} ({d['kind_how']}); {(d.get('ci') or {}).get('source', '')}: "
                 f"страниц {len(d['pages'])}, людей {len(pp)}, на проверку {sum(bool(p['flags']) for p in pp)}"
                 + (f"; пропущены страницы {_ranges(d['gaps'])}" if d["gaps"] else "")
                 + (f"; не загрузились {_ranges([f.get('page') for f in d['failed']])}" if d["failed"] else ""))
    L += ["", "Чаще всего проверить:"] + [f"  {n:5}  {f}" for f, n in flags.most_common(15)]
    L += ["", f"Предложений исправлений ждут решения: {len(ctx.proposals)} (лист «Исправления»)",
          f"Применено исправлений: {ctx.corr.count()}", "",
          f"Файлы: {Path(res['xlsx']).name}, {GED}, {JSONF}"]
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")
