"""
Проверка общего скрипта yaarchive на записанных страницах (без сети):
  6 исповедных ведомостей (1739–1864) + ревизская сказка 1850 (стр. 165) + одна «незагрузившаяся» страница.
Проверяется: разбор обоих типов, тип документа по тексту, GEDCOM (ссылки и формат), навык «Исправление ошибок»
(предложения, решение «да» → правка применяется и переживает пересчёт, «нет» → больше не предлагается,
«отменить» → правка снимается), файл дозагрузки, ручная правка в Excel → предложение на подтверждение.
Запуск: python yaarchive/tests/test_yaarchive.py
"""
import json
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(HERE.parent))
import yaarchive as Y  # noqa: E402
import ya_fix as FX    # noqa: E402

FIX = REPO / "yandex-archive-extension" / "tests" / "fixtures"
REV = REPO / "skill" / "yaarchive-revision-rows" / "assets" / "page165_fixture.json"
ok = True


def check(cond, msg):
    global ok
    print(("OK   " if cond else "ОШИБКА ") + msg)
    ok = ok and bool(cond)


def make_root(tmp):
    root = tmp / "Ссылки"
    pk = root / "yaarchive_debug" / "2026-10-28_10-00-00"
    pk.mkdir(parents=True)
    pages, links = [], []
    for f in sorted(FIX.glob("confession_*_fixture.json")):
        d = json.loads(f.read_text(encoding="utf-8"))
        y, n = int(f.name.split("_")[1]), int(f.name.split("_p")[1].split("_")[0])
        fid = f"{y:04d}aaaa-0000-4000-8000-0000000{n:05d}"
        d["props"]["pageProps"].update(year=y, fileId=fid, mode="confession", docNumber=n)
        (pk / f"{y}_{fid[:8]}_p{n}_next_data.json").write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        pages.append({"year": y, "fileId": fid, "page": n})
        links.append([y, f"https://ya.ru/archive/search?text={y}&sheet_page_from={n}&sheet_page_to={n}&index=archive&fileId={fid}", f"{n}"])
    links.append([1845, "https://ya.ru/archive/search?text=1845&sheet_page_from=892&sheet_page_to=893&index=archive&fileId=1845aaaa-0000-4000-8000-000000000892", "892–893"])
    pages.append({"year": 1845, "fileId": "1845aaaa-0000-4000-8000-000000000892", "page": 893, "error": "таймаут"})
    d = json.loads(REV.read_text(encoding="utf-8"))
    pp = d["props"]["pageProps"] if "props" in d else {"currentNode": d}
    fid = "bfed2966-e671-4c18-bbb0-a9551ad28f8e"
    data = {"props": {"pageProps": {"url": f"https://ya.ru/archive/catalog/{fid}/165", "docNumber": 165, "year": 1850,
                                    "fileId": fid, "mode": "revision", "currentNode": pp.get("currentNode", pp)}}}
    (pk / f"1850_{fid[:8]}_p165_next_data.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    (pk / "run_summary.json").write_text(json.dumps({"version": "2.2.0", "mode": "confession", "complete": True,
                                                     "pages": pages, "links": links}, ensure_ascii=False), encoding="utf-8")
    return root


def run(root):
    Y.main([str(root), "--once", "--no-dialog"])
    out = root / "Итог"
    return json.loads((out / "Итог_люди.json").read_text(encoding="utf-8")), out


def main():
    tmp = Path(tempfile.mkdtemp())
    try:
        root = make_root(tmp)
        data, out = run(root)
        people = data["people"]
        kinds = {d["year"]: d["kind"] for d in data["documents"]}
        check(len(people) == 486, f"людей {len(people)} (458 исповедных + 28 ревизских)")
        check(kinds.get(1850) == "Ревизская сказка" and kinds.get(1739) == "Исповедная ведомость",
              "тип документа: 1850 — ревизская сказка (хотя запуск «исповедные»), 1739 — исповедная")
        rev = [p for p in people if p["year"] == 1850]
        check(len(rev) == 28 and any(p["ref_id"] for p in rev), f"ревизская сказка: {len(rev)} чел., родство по ID")
        conf = [p for p in people if p["kind"] == "confession"]
        check(sum(bool(p["ref_id"]) for p in conf) > 0.4 * len(conf), "исповедные: родство (жена/дети) у большинства")
        check(all(isinstance(p["birth"], int) for p in conf if re.fullmatch(r"\d+", str(p["age"]))), "год рождения у всех с возрастом")
        check(not any(p["settlement"].startswith("д. ") for p in people), "НП к одному виду («деревня …»)")
        # GEDCOM
        L = (out / "Итог_древо.ged").read_text(encoding="utf-8").splitlines()
        defs = {m.group(1) for m in (re.match(r"^0 (@\w+@)", x) for x in L) if m}
        refs = [x.split()[-1] for x in L if re.match(r"^\d (FAMS|FAMC|HUSB|WIFE|CHIL|SOUR) @", x)]
        check(L[0] == "0 HEAD" and L[-1] == "0 TRLR" and all(r in defs for r in refs) and max(map(len, L)) <= 255,
              f"GEDCOM: {sum(x.endswith(' INDI') for x in L)} персон, {sum(x.endswith(' FAM') for x in L)} семей, ссылки целы")
        # Исправления
        props = json.loads((out / ".предложения.json").read_text(encoding="utf-8"))
        types = {p["type"] for p in props.values()}
        check("Дозагрузка страниц" in types, "предложена дозагрузка (стр. 893 не загрузилась)")
        check(len(props) >= 5, f"предложений {len(props)}: {sorted(types)}")
        corr = FX.Corrections(out)
        yes = next(p for p in props.values() if p["changes"])
        no = next(p for p in props.values() if p["changes"] and p["id"] != yes["id"])
        dl = next(p for p in props.values() if p["type"] == "Дозагрузка страниц")
        corr.decide(yes, "да", "тест", roots=[root])
        corr.decide(no, "нет", "тест", roots=[root])
        corr.decide(dl, "да", "тест", roots=[root])
        data2, _ = run(root)
        key, field, _old, new = yes["changes"][0]
        p2 = next(p for p in data2["people"] if p["key"] == key)
        check(str(p2[field]) == str(new) and p2["fixed"], f"«да»: {field} → {new} применено после пересчёта")
        props2 = json.loads((out / ".предложения.json").read_text(encoding="utf-8"))
        check(no["id"] not in props2 and yes["id"] not in props2, "решённые больше не предлагаются")
        check((root / "Дозагрузка_страниц.html").exists() and "sheet_page_from=893" in (root / "Дозагрузка_страниц.html").read_text(encoding="utf-8"),
              "файл дозагрузки создан со стр. 893")
        # ручная правка в Excel
        from openpyxl import load_workbook
        time.sleep(3)
        x = out / "Итог_люди.xlsx"
        wb = load_workbook(x)
        ws = wb["Люди"]
        head = [c.value for c in ws[1]]
        ws.cell(2, head.index("Возраст") + 1).value = 77
        wb.save(x)
        snap = json.loads((out / ".снимок.json").read_text(encoding="utf-8"))
        snap["written"] -= 10
        (out / ".снимок.json").write_text(json.dumps(snap, ensure_ascii=False), encoding="utf-8")
        import os
        os.utime(x, (time.time() - 5, time.time() - 5))
        fb = FX.read_excel_feedback(out)
        man = FX.load_manual(out)
        check(fb and any(p["new"] == "77" for p in man.values()), f"ручная правка найдена и ждёт подтверждения ({fb})")
        FX.Corrections(out).undo(yes["id"])
        data3, _ = run(root)
        p3 = next(p for p in data3["people"] if p["key"] == key)
        check(not p3["fixed"] or all(field not in f for f in p3["fixed"]), "«отменить»: правка снята")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("ИТОГ:", "OK" if ok else "ЕСТЬ ОШИБКИ")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
