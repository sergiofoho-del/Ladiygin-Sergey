#!/usr/bin/env python3
"""
Проставить в result.xlsx ссылки на уже сохранённые сканы (столбец «Локальная Ссылка на страницу»).

Имя файла скана однозначно следует из столбца «Архив хранения документа»
(«ЦГА Москвы/Фонд 203/Опись 747/Дело 15/Скан 865» (или с «Главная/Каталог/» впереди) -> «ЦГА_Москвы_[ЯА]_203-747-15-0865.jpg»),
так же называет файлы Scrpt_YA_Orig_Scan и навык «Сохранить Скан». Поэтому сканы, скачанные отдельно
(userscript, ya_orig_scan.js, расширение), подключаются к готовому Excel без повторной обработки.

Запуск: python link_scans.py result.xlsx [--scans /Volumes/DATA/Archive_Scans] [--out result_links.xlsx]
В столбце — локальный путь папки сканов + имя файла (по умолчанию /Volumes/DATA/Archive_Scans/<файл>.jpg), ссылка — file://… (как в расширении 2.8.0).
"""
import argparse
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

HEAD_URL, HEAD_LOCAL, HEAD_CGA = "Ссылка на страницу", "Локальная Ссылка на страницу", "Архив хранения документа"


def scan_name(cga):
    """«Главная/Каталог/ЦГА Москвы/Фонд 203/Опись 747/Дело 15/Скан 865» -> имя файла скана"""
    parts = [p.strip() for p in str(cga or "").split("/")]
    if parts[:2] == ["Главная", "Каталог"]:          # прежний вид (до расширения 2.3.0) — с «Главная/Каталог»
        parts = parts[2:]
    if len(parts) < 3:
        return None
    d = {}
    for p in parts:
        m = re.match(r"(Фонд|Опись|Дело|Скан)\s+(.+)$", p)
        if m:
            d[m.group(1)] = m.group(2).strip()
    archive = next((p for p in parts if not re.match(r"(Фонд|Опись|Дело|Скан)\s", p)), "")
    if "Скан" not in d or not d["Скан"].isdigit():
        return None
    codes = [d.get("Фонд", ""), d.get("Опись", ""), d.get("Дело", "")]
    codes = [c for c in codes if c]
    if len(codes) < 2:
        return None
    clean = lambda v: re.sub(r"_+", "_", re.sub(r"\s+", "_", re.sub(r'[<>:"/\\|?*\x00-\x1F]', "", v))).strip("_")
    prefix = clean(archive) + "_[ЯА]_" if archive else ""
    return prefix + clean("-".join(codes) + "-" + d["Скан"].zfill(4)) + ".jpg"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--scans", default="/Volumes/DATA/Archive_Scans", help="папка со сканами (Archive_Scans); по умолчанию /Volumes/DATA/Archive_Scans")
    ap.add_argument("--out", default=None, help="куда сохранить (по умолчанию — поверх исходного файла)")
    a = ap.parse_args()
    xlsx = Path(a.xlsx).resolve()
    scans = Path(a.scans) if Path(a.scans).is_absolute() else (xlsx.parent / a.scans)
    scans = scans.resolve()
    have = {p.name for p in scans.glob("*.jpg")} if scans.is_dir() else set()
    wb = load_workbook(xlsx)
    linked = missing = 0
    for ws in wb.worksheets:
        head = {c.value: c.column for c in ws[1] if c.value}
        if HEAD_CGA not in head and "Ссылка на ЦГА Москвы" in head:     # файлы до версии 2.8.0
            head[HEAD_CGA] = head["Ссылка на ЦГА Москвы"]
        if HEAD_CGA not in head:
            continue
        if HEAD_LOCAL not in head:              # старый файл без столбца — добавляем в конец
            col = ws.max_column + 1
            ws.cell(1, col, HEAD_LOCAL).font = Font(bold=True)
            head[HEAD_LOCAL] = col
        for r in range(2, ws.max_row + 1):
            cga = ws.cell(r, head[HEAD_CGA]).value
            name = scan_name(cga)
            if not name:
                continue
            if name in have:
                full = (scans / name).as_posix()          # локальный путь + имя файла
                cell = ws.cell(r, head[HEAD_LOCAL], full)
                cell.hyperlink = scans.joinpath(name).as_uri()      # file:///…
                cell.font = Font(color="0563C1", underline="single")
                linked += 1
            else:
                missing += 1
    out = a.out or str(xlsx)
    wb.save(out)
    print(f"ссылок на сканы: {linked}, сканов нет в папке: {missing}; папка {scans} ({len(have)} файлов) -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
